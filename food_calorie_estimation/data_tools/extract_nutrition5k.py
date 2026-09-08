"""
Extract Nutrition5k Kaggle dataset (pkl + xlsx).
Reads dish_images.pkl, saves RGB images as jpg, and creates nutrition labels CSV.

Usage:
    python extract_nutrition5k.py --pkl_path <path to dish_images.pkl> --xlsx_dir <path to xlsx dir> --out_dir <output dir>
"""
import os
import sys
import pickle
import argparse
import struct
import io
from pathlib import Path

def extract_images_from_pkl(pkl_path, out_dir, max_dishes=None):
    """Extract RGB images from dish_images.pkl and save as jpg."""
    img_dir = os.path.join(out_dir, "images")
    os.makedirs(img_dir, exist_ok=True)
    
    print(f"Loading pkl from {pkl_path} ...")
    print("(This may take 1-2 minutes for 2.5GB file)")
    
    with open(pkl_path, "rb") as f:
        data = pickle.load(f)
    
    # Determine data structure
    print(f"PKL type: {type(data)}")
    
    if isinstance(data, dict):
        dish_ids = list(data.keys())
        print(f"Found {len(dish_ids)} dishes (dict format)")
    elif isinstance(data, list):
        dish_ids = range(len(data))
        print(f"Found {len(dish_ids)} dishes (list format)")
    else:
        print(f"Unexpected data type: {type(data)}")
        print(f"Data: {data}")
        return
    
    if max_dishes:
        dish_ids = dish_ids[:max_dishes]
    
    count = 0
    skipped = 0
    for i, did in enumerate(dish_ids):
        entry = data[did] if isinstance(data, dict) else data[i]
        
        try:
            # entry could be dict with 'rgb_image'/'depth_image' keys
            # or tuple (dish_id, rgb_bytes, depth_bytes)
            rgb_bytes = None
            
            if isinstance(entry, dict):
                # Try common key names
                for key in ['rgb_image', 'rgb', 'image', 'img']:
                    if key in entry:
                        rgb_bytes = entry[key]
                        break
                dish_name = entry.get('dish_id', did)
            elif isinstance(entry, (list, tuple)):
                # Usually (dish_id, rgb_bytes, depth_bytes)
                if len(entry) >= 2:
                    dish_name = entry[0] if entry[0] else did
                    rgb_bytes = entry[1]
                else:
                    rgb_bytes = entry[0]
                    dish_name = did
            else:
                rgb_bytes = entry
                dish_name = did
            
            if rgb_bytes is None:
                skipped += 1
                continue
            
            # Save image
            img_filename = f"{dish_name}.jpg" if dish_name else f"dish_{i:05d}.jpg"
            img_path = os.path.join(img_dir, img_filename)
            
            if isinstance(rgb_bytes, bytes):
                with open(img_path, "wb") as img_f:
                    img_f.write(rgb_bytes)
            else:
                # Might be a PIL Image or numpy array
                try:
                    from PIL import Image
                    import numpy as np
                    if isinstance(rgb_bytes, Image.Image):
                        rgb_bytes.save(img_path)
                    elif isinstance(rgb_bytes, np.ndarray):
                        Image.fromarray(rgb_bytes).save(img_path)
                    else:
                        print(f"  Unknown image type for dish {did}: {type(rgb_bytes)}")
                        skipped += 1
                        continue
                except ImportError:
                    print(f"  Need PIL/numpy for non-bytes images. Skipping dish {did}")
                    skipped += 1
                    continue
            
            count += 1
            if count % 500 == 0:
                print(f"  Extracted {count} images...")
                
        except Exception as e:
            skipped += 1
            if skipped <= 5:
                print(f"  Error on dish {did}: {e}")
            continue
    
    print(f"\nExtraction complete: {count} images saved, {skipped} skipped")
    print(f"Images saved to: {img_dir}")
    return img_dir


def parse_xlsx_labels(xlsx_dir, out_dir):
    """Read xlsx nutrition files and create summary CSV."""
    import csv
    
    try:
        import openpyxl
    except ImportError:
        print("Installing openpyxl...")
        os.system(f"{sys.executable} -m pip install openpyxl -q")
        import openpyxl
    
    csv_path = os.path.join(out_dir, "nutrition_labels.csv")
    
    # Read dishes.xlsx for dish-level nutrition
    dishes_path = os.path.join(xlsx_dir, "dishes.xlsx")
    if not os.path.exists(dishes_path):
        print(f"dishes.xlsx not found at {dishes_path}")
        return
    
    print(f"\nReading {dishes_path} ...")
    wb = openpyxl.load_workbook(dishes_path, read_only=True)
    ws = wb.active
    
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    print(f"Columns: {header}")
    
    with open(csv_path, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for row in rows[1:]:
            writer.writerow(row)
    
    wb.close()
    print(f"Saved {len(rows)-1} dish labels to {csv_path}")
    
    # Also read dish_ingredients.xlsx if exists
    ingr_path = os.path.join(xlsx_dir, "dish_ingredients.xlsx")
    if os.path.exists(ingr_path):
        ingr_csv = os.path.join(out_dir, "dish_ingredients.csv")
        print(f"\nReading {ingr_path} ...")
        wb2 = openpyxl.load_workbook(ingr_path, read_only=True)
        ws2 = wb2.active
        rows2 = list(ws2.iter_rows(values_only=True))
        header2 = rows2[0]
        with open(ingr_csv, "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(header2)
            for row in rows2[1:]:
                writer.writerow(row)
        wb2.close()
        print(f"Saved {len(rows2)-1} ingredient rows to {ingr_csv}")


def main():
    parser = argparse.ArgumentParser(description="Extract Nutrition5k Kaggle dataset")
    parser.add_argument("--pkl_path", required=True, help="Path to dish_images.pkl")
    parser.add_argument("--xlsx_dir", required=True, help="Directory containing xlsx files")
    parser.add_argument("--out_dir", required=True, help="Output directory")
    parser.add_argument("--max_dishes", type=int, default=None, help="Max dishes to extract (for testing)")
    args = parser.parse_args()
    
    os.makedirs(args.out_dir, exist_ok=True)
    
    # Step 1: Extract images from pkl
    if os.path.exists(args.pkl_path):
        extract_images_from_pkl(args.pkl_path, args.out_dir, args.max_dishes)
    else:
        print(f"PKL not found: {args.pkl_path}")
    
    # Step 2: Parse xlsx labels
    parse_xlsx_labels(args.xlsx_dir, args.out_dir)
    
    print("\n=== Done ===")


if __name__ == "__main__":
    main()
