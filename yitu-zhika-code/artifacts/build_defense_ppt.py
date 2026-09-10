#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""从 artifacts/ppt_defense_content.json 生成答辩 PPT。

依赖仓库内 vendored 的 python-pptx（.pptx_libs，为 cp314 编译）：

    set PYTHONPATH=<repo>\\.pptx_libs
    C:\\Users\\user\\AppData\\Local\\Python\\pythoncore-3.14-64\\python.exe artifacts\\build_defense_ppt.py

坐标系：JSON 使用 960 x 540 pt 的画布（16:9），本脚本按 1 pt = 12700 EMU 换算。
"""
from __future__ import annotations

import json
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

EMU_PER_PT = 12700

TITLE_RGB = RGBColor(0x1F, 0x38, 0x64)
ACCENT_RGB = RGBColor(0x2E, 0x7D, 0x32)
BODY_RGB = RGBColor(0x26, 0x26, 0x26)
MUTED_RGB = RGBColor(0x78, 0x78, 0x78)
BAND_RGB = RGBColor(0xF0, 0xF8, 0xF0)
WHITE_RGB = RGBColor(0xFF, 0xFF, 0xFF)
RED_RGB = RGBColor(0xB2, 0x22, 0x22)
PIC_BORDER_RGB = RGBColor(0xD7, 0xDF, 0xE8)

TABLE_STYLE_GRID = "{5940675A-B579-460E-94D1-54222C63F5DA}"

BULLET_0 = "\u25a0  "   # filled square
BULLET_1 = "\u25cb  "   # hollow circle

COLOR_BY_NAME = {"accent": ACCENT_RGB, "red": RED_RGB}


def pt(value: float) -> Emu:
    return Emu(int(round(value * EMU_PER_PT)))


def set_run(run, font: str, size: float, bold: bool, color: RGBColor) -> None:
    run.font.size = Pt(size)
    run.font.bold = bool(bold)
    run.font.color.rgb = color
    run.font.name = font
    # 让中日韩文字也走同一字体，否则会回退到宋体
    rPr = run._r.get_or_add_rPr()
    for getter in ("get_or_add_ea", "get_or_add_cs"):
        try:
            el = getattr(rPr, getter)()
            el.set("typeface", font)
        except Exception:  # noqa: BLE001
            pass


def add_textbox(slide, left: float, top: float, width: float, height: float):
    box = slide.shapes.add_textbox(pt(left), pt(top), pt(width), pt(height))
    tf = box.text_frame
    tf.word_wrap = True
    return box, tf


def style_paragraph(par, font, size, bold, color, level=0, space_after=6.0):
    par.level = level
    par.space_after = Pt(space_after)
    par.space_before = Pt(0)
    for run in par.runs:
        set_run(run, font, size, bold, color)


def add_accent_bar(slide, left, top, width, height):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, pt(left), pt(top), pt(width), pt(height))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_RGB
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def add_table(slide, spec, font: str):
    left, top = float(spec["left"]), float(spec["top"])
    width, height = float(spec["width"]), float(spec["height"])
    headers = [str(h) for h in spec["headers"]]
    rows = [[str(c) for c in r] for r in spec["rows"]]
    n_cols = len(headers)
    n_rows = len(rows) + 1

    size = float(spec.get("font_size") or 11.0)
    if n_rows > 9:
        size -= 1.0
    size = max(size, 8.0)

    shape = slide.shapes.add_table(n_rows, n_cols, pt(left), pt(top), pt(width), pt(height))
    table = shape.table
    table.first_row = True
    table.horz_banding = False

    tblPr = table._tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        for el in tblPr.findall(qn("a:tableStyleId")):
            tblPr.remove(el)
        style_el = etree.SubElement(tblPr, qn("a:tableStyleId"))
        style_el.text = TABLE_STYLE_GRID

    col_w = width / n_cols
    for idx in range(n_cols):
        table.columns[idx].width = pt(col_w)
    row_h = height / n_rows
    for idx in range(n_rows):
        table.rows[idx].height = pt(row_h)

    def fill_cell(cell, text, bold, color, fill_rgb, align_center=True):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill_rgb
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = pt(3)
        cell.margin_right = pt(3)
        cell.margin_top = pt(1.5)
        cell.margin_bottom = pt(1.5)
        tf = cell.text_frame
        tf.word_wrap = True
        par = tf.paragraphs[0]
        par.alignment = PP_ALIGN.CENTER if align_center else PP_ALIGN.LEFT
        run = par.add_run()
        run.text = text
        set_run(run, font, size, bold, color)

    for c, header in enumerate(headers):
        fill_cell(table.cell(0, c), header, True, WHITE_RGB, ACCENT_RGB)

    for r, row in enumerate(rows):
        for c in range(n_cols):
            value = row[c] if c < len(row) else ""
            fill = BAND_RGB if r % 2 == 1 else WHITE_RGB
            fill_cell(table.cell(r + 1, c), value, c == 0, BODY_RGB, fill)
    return shape


def add_picture(slide, path: str, spec, base_dir: str):
    left, top = float(spec["left"]), float(spec["top"])
    width, height = float(spec["width"]), float(spec["height"])
    with Image.open(path) as im:
        img_w, img_h = im.size
    scale = min(width / img_w, height / img_h)
    fit_w, fit_h = img_w * scale, img_h * scale
    x = left + (width - fit_w) / 2.0
    y = top + (height - fit_h) / 2.0
    pic = slide.shapes.add_picture(
        path, pt(x), pt(y), width=pt(fit_w), height=pt(fit_h)
    )
    try:
        pic.line.color.rgb = PIC_BORDER_RGB
        pic.line.width = Pt(0.75)
    except Exception:  # noqa: BLE001
        pass
    return pic


def build(cfg: dict, base_dir: str, out_path: str) -> int:
    font = cfg.get("font") or "Microsoft YaHei"
    footer = cfg.get("footer") or ""
    slides_spec = cfg["slides"]
    total = len(slides_spec)

    prs = Presentation()
    prs.slide_width = pt(960)
    prs.slide_height = pt(540)
    blank = prs.slide_layouts[6]

    deferred = []
    for index, spec in enumerate(slides_spec, start=1):
        slide = prs.slides.add_slide(blank)
        title = spec.get("title") or ""

        if title:
            _, tf = add_textbox(slide, 30, 16, 900, 46)
            tf.text = title
            style_paragraph(tf.paragraphs[0], font, 23, True, TITLE_RGB, 0, 0)
            add_accent_bar(slide, 30, 66, 900, 3)

        for sh in spec["shapes"]:
            kind = sh["type"]
            if kind == "caption":
                _, tf = add_textbox(slide, sh["left"], sh["top"], sh["width"], sh["height"])
                tf.text = sh["text"]
                size = float(sh.get("font_size") or 11.0)
                color = TITLE_RGB if size >= 30 else BODY_RGB
                align = PP_ALIGN.CENTER if sh.get("align") == "center" else PP_ALIGN.LEFT
                for par in tf.paragraphs:
                    par.alignment = align
                    style_paragraph(par, font, size, bool(sh.get("bold")), color, 0, 0)
            elif kind == "bullets":
                _, tf = add_textbox(slide, sh["left"], sh["top"], sh["width"], sh["height"])
                items = sh["items"]
                lines = []
                for it in items:
                    lvl = int(it.get("level") or 0)
                    lines.append((BULLET_0 if lvl == 0 else BULLET_1) + it["text"])
                tf.text = "\n".join(lines)
                for par, it in zip(tf.paragraphs, items):
                    lvl = int(it.get("level") or 0)
                    size = 13.0 if lvl == 0 else 11.5
                    color = COLOR_BY_NAME.get(it.get("color"), BODY_RGB)
                    style_paragraph(par, font, size, bool(it.get("bold")), color, lvl, 6.0)
            elif kind == "table":
                add_table(slide, sh, font)
            elif kind == "picture":
                img = os.path.join(base_dir, sh["path"].replace("/", os.sep))
                if os.path.isfile(img):
                    add_picture(slide, img, sh, base_dir)
                else:
                    _, tf = add_textbox(slide, sh["left"], sh["top"], sh["width"], sh["height"])
                    tf.text = "[missing image] " + sh["path"]
                    style_paragraph(tf.paragraphs[0], font, 10, False, RED_RGB, 0, 0)
                    print("WARN missing image: %s" % img)
            else:
                raise SystemExit("unknown shape type: %s" % kind)

        if not title:
            add_accent_bar(slide, 360, 224, 240, 5)

        if title and footer:
            _, tf = add_textbox(slide, 30, 508, 780, 20)
            tf.text = footer
            style_paragraph(tf.paragraphs[0], font, 9, False, MUTED_RGB, 0, 0)
            _, tf2 = add_textbox(slide, 820, 508, 110, 20)
            tf2.text = "%d / %d" % (index, total)
            par = tf2.paragraphs[0]
            par.alignment = PP_ALIGN.RIGHT
            style_paragraph(par, font, 9, False, MUTED_RGB, 0, 0)

        notes = spec.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

        deferred.append(index)

    prs.save(out_path)
    return len(deferred)


def main() -> int:
    here = os.path.dirname(os.path.abspath(__file__))
    base_dir = os.path.dirname(here)
    json_path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "ppt_defense_content.json")
    with open(json_path, "r", encoding="utf-8") as fh:
        cfg = json.load(fh)
    out_path = os.path.join(here, cfg["out_file"])
    if os.path.exists(out_path):
        os.remove(out_path)
    count = build(cfg, base_dir, out_path)
    print("OK slides=%d" % count)
    print("FILE %s" % out_path)
    print("SIZE %d bytes" % os.path.getsize(out_path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
